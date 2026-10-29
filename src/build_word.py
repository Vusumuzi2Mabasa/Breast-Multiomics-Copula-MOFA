"""Render the shared manuscript with Quarto, then apply stable Word styles."""
from pathlib import Path
import argparse,json,os,re,shutil,subprocess
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]


def reference(path):
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.7)
    s.left_margin=s.right_margin=Inches(.75)
    for style in d.styles:
        if style.type==1:
            style.font.name='Calibri';style.font.size=Pt(11);style.font.color.rgb=RGBColor(0,0,0)
            style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.04
    for name,size in [('Title',23),('Subtitle',12),('Heading 1',16),('Heading 2',13),('Heading 3',11)]:
        st=d.styles[name];st.font.size=Pt(size);st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(10)
    for e in list(d.styles.element.iter(qn('w:pBdr'))):e.getparent().remove(e)
    d.save(path)


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--quarto',default=os.environ.get('QUARTO') or shutil.which('quarto'))
    args=ap.parse_args()
    if not args.quarto:ap.error('Quarto is required; set QUARTO or pass --quarto.')
    work=ROOT/'work/word';work.mkdir(parents=True,exist_ok=True)
    reference(work/'reference.docx')
    readme=(ROOT/'README.md').read_text(encoding='utf-8')
    content=readme.split('<!-- manuscript -->',1)[1].lstrip()
    # Keep a scientific report; detailed machine commands remain in workflow.txt.
    content=content.split('## Reproducibility',1)[0]+ '\n## References'+readme.split('## References',1)[1]
    disclosure=next(line for line in readme.splitlines() if line.startswith('**Simulation study.**'))
    content=disclosure+'\n\n'+content
    content=content.replace('](figures/','](../figures/')
    repo='https://github.com/Vusumuzi2Mabasa/Breast-Multiomics-Copula-MOFA/blob/main/'
    content=re.sub(r'(?<!!)\[([^]]+)\]\((?!https?://|#)([^)]+)\)',lambda m:f'[{m[1]}]({repo}{m[2]})',content)
    header='''---
title: "Copula and MOFA Integration of Simulated Blood Multiomics"
author: "Vusumuzi Mabasa"
subtitle: "A reproducible simulation benchmark"
format:
  docx:
    reference-doc: ../work/word/reference.docx
    toc: false
    number-sections: false
execute:
  enabled: false
---

'''
    qmd=ROOT/'report/manuscript.qmd';qmd.write_text(header+content,encoding='utf-8')
    subprocess.run([args.quarto,'render',str(qmd),'--to','docx','--output','Copula_MOFA_Simulation_Report.docx'],cwd=ROOT,check=True)
    out=ROOT/'report/Copula_MOFA_Simulation_Report.docx';d=Document(out)
    d.core_properties.author='Vusumuzi Mabasa'
    for st in d.styles:
        if st.type==1:st.font.color.rgb=RGBColor(0,0,0)
    for p in d.paragraphs:
        if p.style.name.startswith('Heading'):p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//w:drawing'):
            p.alignment=1;p.paragraph_format.keep_with_next=True
        if p.text.startswith('Figure '):p.paragraph_format.space_after=Pt(9)
        if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
    for shape in d.inline_shapes:
        ratio=Inches(6.4)/shape.width;shape.width=Inches(6.4);shape.height=int(shape.height*ratio)
    for table in d.tables:
        table.autofit=True
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side)
            for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
            borders.append(e)
        table._tbl.tblPr.append(borders)
        for i,row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for cell in row.cells:
                cell.vertical_alignment=1
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF4' if i==0 else ('F5F7F9' if i%2 else 'FFFFFF'))
                cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                    if len(table.rows)<=8:p.paragraph_format.keep_with_next=True
                    for r in p.runs:r.font.size=Pt(9.5);r.bold=i==0
    for st in d.styles:
        for b in list(st.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
    for s in d.sections:
        p=s.footer.paragraphs[0];p.alignment=2
        p.add_run('Vusumuzi Mabasa | Simulated data | ')
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
        for r in p.runs:r.font.size=Pt(8)
    d.save(out)
    equations=len(d.element.xpath('.//m:oMath'))
    audit=dict(author=d.core_properties.author,figures=len(d.inline_shapes),tables=len(d.tables),editable_equations=equations)
    (ROOT/'results/document_structure.json').write_text(json.dumps(audit,indent=2));print(json.dumps(audit))


if __name__=='__main__':main()
