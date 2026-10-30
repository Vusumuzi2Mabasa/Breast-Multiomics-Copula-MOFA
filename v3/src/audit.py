"""Audit new results while checking byte preservation of both earlier studies."""
from pathlib import Path
import hashlib,json,argparse
from docx import Document
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT.parent


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()


def main():
    ap=argparse.ArgumentParser();ap.add_argument('--fresh',action='store_true');a=ap.parse_args()
    lock=json.loads((ROOT/'protocol/source_lock.json').read_text())
    for path,digest in lock['sha256'].items():assert sha(ROOT/path)==digest,path
    preserved=0
    if not a.fresh:
        old=json.loads((ROOT/'protocol/preservation.json').read_text())
        for path,digest in old['sha256'].items():assert sha(BASE/path)==digest,path;preserved+=1
    num=json.loads((ROOT/'results/numerical_audit.json').read_text());assert num['status']=='passed' and not num['failures']
    assert num['eqtl_cohorts']==1600 and num['eqtl_gene_records']==19200 and num['integration_cohorts']==80 and num['integration_method_records']==560
    assert json.loads((ROOT/'results/R_validation.json').read_text())['status']=='passed'
    assert json.loads((ROOT/'results/integration_replay.json').read_text())['status']=='passed'
    amendment=json.loads((ROOT/'protocol/robustness_amendment.json').read_text())
    assert sha(ROOT/'src/robustness.py')==amendment['source_sha256']
    assert json.loads((ROOT/'results/robustness_audit.json').read_text())['fresh_cohorts']==800
    assert json.loads((ROOT/'results/R_HC3_validation.json').read_text())['status']=='passed'
    assert json.loads((ROOT/'results/example_QC.json').read_text())['status']=='passed'
    cal_seeds=set();train_seeds=set();test_seeds=set()
    for file in (ROOT/'results/eqtl').glob('*.json'):
        d=json.loads(file.read_text());assert d['status']=='ok' and d['seed'] not in cal_seeds;cal_seeds.add(d['seed'])
    for file in (ROOT/'results/integration').glob('*.json'):
        d=json.loads(file.read_text());assert d['status']=='ok' and d['seed'] not in train_seeds and d['test_seed'] not in test_seeds
        train_seeds.add(d['seed']);test_seeds.add(d['test_seed'])
        r=json.loads((ROOT/f'results/diablo/{file.name}').read_text());assert r['status']=='ok'
    assert not cal_seeds&train_seeds and not cal_seeds&test_seeds and not train_seeds&test_seeds
    fresh_seeds=set()
    for file in (ROOT/'results/robustness').glob('fresh_*.json'):
        d=json.loads(file.read_text());assert d['seed'] not in fresh_seeds;fresh_seeds.add(d['seed'])
    assert len(fresh_seeds)==800 and not fresh_seeds&(cal_seeds|train_seeds|test_seeds)
    assert len(list((ROOT/'figures').glob('*.png')))==7
    doc=Document(ROOT/'report/Genotype_RNA_eQTL_Report.docx')
    assert doc.core_properties.author=='Vusumuzi Mabasa' and len(doc.inline_shapes)==7 and len(doc.tables)==7
    assert [len(t.rows) for t in doc.tables]==[6,8,5,5,9,5,13]
    assert len(doc.element.xpath('.//m:oMath'))>=10
    readme=(BASE/'README.md').read_text(encoding='utf-8');assert 'nan%' not in readme.lower() and 'inf%' not in readme.lower()
    assert 'Protein, somatic alteration and copula models are excluded' in readme
    if not a.fresh:
        qa=json.loads((ROOT/'results/report_qa.json').read_text());assert qa['status']=='passed'
        assert sha(ROOT/'report/Genotype_RNA_eQTL_Report.docx')==qa['word_sha256']
        assert sha(BASE/'README.md')==qa['readme_sha256']
    result=dict(status='passed',preserved_files=preserved,eqtl_cohorts=1600,fresh_null_cohorts=800,eqtl_gene_records=19200,integration_pairs=80,diablo_fits=80,method_records=560,diagnostic_flags=num['diagnostic_flags'],word_pages=None if a.fresh else qa['pages'],eqtl_calibration='Negative-control excesses observed; neither ordinary F nor HC3 certified for confirmatory use without further calibration.')
    if not a.fresh:(ROOT/'results/release_audit.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()
